# ============================================================
#  DATA CLEANING & REPORTING AUTOMATION (Enterprise Edition)
# ============================================================

import pandas as pd
import numpy as np
import os
import warnings
import logging
from datetime import datetime
from typing import List, Dict, Any

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

# Configure professional logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S"
)

# ──────────────────────────────────────────────────────────────
#  ⚙️  CONFIGURATION
# ──────────────────────────────────────────────────────────────

INPUT_FILE       = "data/raw/sales_data.csv"
OUTPUT_CSV       = "data/cleaned/sales_data_cleaned.csv"
OUTPUT_REPORT    = "reports/cleaning_report.xlsx"

NUMERIC_COLUMNS  = ["Sales Amount", "Quantity"]
DATE_COLUMNS     = ["Order Date"]
TEXT_COLUMNS     = ["Customer Name", "City", "Category"]
CATEGORY_COLUMN  = "Category"
VALUE_COLUMN     = "Sales Amount"


# ──────────────────────────────────────────────────────────────
#  DATA PROCESSING FUNCTIONS
# ──────────────────────────────────────────────────────────────

def load_data(filepath: str) -> pd.DataFrame:
    """Load a CSV or Excel file into a pandas DataFrame."""
    logging.info(f"Loading data from: {filepath}")
    
    if not os.path.exists(filepath):
        logging.error(f"File not found: {filepath}")
        raise FileNotFoundError(f"Could not find file: {filepath}")
        
    try:
        if filepath.endswith(".csv"):
            df = pd.read_csv(filepath)
        elif filepath.endswith((".xlsx", ".xls")):
            df = pd.read_excel(filepath)
        else:
            raise ValueError("File must be .csv, .xlsx, or .xls")
            
        logging.info(f"Successfully loaded {len(df):,} rows and {len(df.columns)} columns")
        return df
    except Exception as e:
        logging.error(f"Failed to load data: {e}")
        raise

def inspect_data(df: pd.DataFrame) -> Dict[str, Any]:
    """Scan the data and report problems."""
    logging.info("Running Data Quality Check...")
    problems = {}
    
    # Missing values
    missing = df.isnull().sum()
    missing = missing[missing > 0]
    if not missing.empty:
        for col, count in missing.items():
            problems[f"Missing — {col}"] = int(count)
            logging.warning(f"Found {count} missing values in '{col}'")
    else:
        logging.info("No missing values found.")
        
    # Duplicate rows
    dup_count = int(df.duplicated().sum())
    if dup_count > 0:
        problems["Duplicate rows"] = dup_count
        logging.warning(f"Found {dup_count} duplicate rows.")
    else:
        logging.info("No duplicate rows found.")
        
    # Text inconsistencies
    for col in df.columns:
        if df[col].dtype == 'object':
            unique_vals = df[col].dropna().unique()
            lowered = [str(v).lower().strip() for v in unique_vals]
            if len(lowered) != len(set(lowered)):
                problems[f"Inconsistent text — {col}"] = "Yes"
                logging.warning(f"Inconsistent casing/spacing detected in '{col}'")
                
    return problems

def fix_column_names(df: pd.DataFrame) -> pd.DataFrame:
    """Standardise column names to Title Case with no trailing spaces."""
    df.columns = [str(col).strip().title() for col in df.columns]
    return df

def convert_numeric_columns(df: pd.DataFrame, numeric_cols: List[str]) -> pd.DataFrame:
    """Force specified columns to numeric type."""
    for col in numeric_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    return df

def remove_duplicates(df: pd.DataFrame) -> pd.DataFrame:
    """Remove exact duplicate rows."""
    before = len(df)
    df = df.drop_duplicates()
    removed = before - len(df)
    if removed > 0:
        logging.info(f"Removed {removed} duplicate rows.")
    return df

def clean_text_columns(df: pd.DataFrame, text_cols: List[str]) -> pd.DataFrame:
    """Clean text by stripping spaces and converting to Title Case."""
    cleaned_cols = []
    for col in text_cols:
        if col in df.columns:
            df[col] = (df[col]
                       .astype(str)
                       .str.strip()
                       .str.replace(r"\s+", " ", regex=True)
                       .str.title())
            df[col] = df[col].replace("Nan", np.nan)
            cleaned_cols.append(col)
    if cleaned_cols:
        logging.info(f"Cleaned text columns: {cleaned_cols}")
    return df

def standardise_categories(df: pd.DataFrame, category_col: str) -> pd.DataFrame:
    """Standardise capitalisation for a categorical column."""
    if category_col in df.columns:
        df[category_col] = df[category_col].astype(str).str.title().str.strip()
        df[category_col] = df[category_col].replace("Nan", np.nan)
        logging.info(f"Standardised categories in '{category_col}'")
    return df

def fill_missing_values(df: pd.DataFrame, numeric_cols: List[str], text_cols: List[str]) -> pd.DataFrame:
    """Fill missing numeric values with median, and text with 'Unknown'."""
    for col in numeric_cols:
        if col in df.columns:
            n_missing = df[col].isnull().sum()
            if n_missing > 0:
                median_val = df[col].median()
                df[col] = df[col].fillna(median_val)
                logging.info(f"Filled {n_missing} blanks in '{col}' with median: {median_val:.2f}")

    for col in text_cols:
        if col in df.columns:
            n_missing = df[col].isnull().sum()
            if n_missing > 0:
                df[col] = df[col].fillna("Unknown")
                logging.info(f"Filled {n_missing} blanks in '{col}' with 'Unknown'")
    return df

def parse_dates(df: pd.DataFrame, date_cols: List[str]) -> pd.DataFrame:
    """Parse string dates into datetime objects."""
    for col in date_cols:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors="coerce")
            bad = df[col].isnull().sum()
            if bad > 0:
                logging.warning(f"Failed to parse {bad} dates in '{col}'")
            else:
                logging.info(f"Successfully parsed dates in '{col}'")
    return df

def flag_outliers(df: pd.DataFrame, value_col: str) -> pd.DataFrame:
    """Flag statistical outliers using the IQR method."""
    if value_col not in df.columns or df[value_col].isnull().all():
        return df

    Q1 = df[value_col].quantile(0.25)
    Q3 = df[value_col].quantile(0.75)
    IQR = Q3 - Q1
    lower = Q1 - 1.5 * IQR
    upper = Q3 + 1.5 * IQR

    flag_col = f"{value_col} Is Outlier"
    df[flag_col] = ~df[value_col].between(lower, upper)
    
    n_outliers = df[flag_col].sum()
    if n_outliers > 0:
        logging.warning(f"Flagged {n_outliers} outliers in '{value_col}'")
    return df

def add_derived_columns(df: pd.DataFrame, date_col: str = "Order Date") -> pd.DataFrame:
    """Add Month, Quarter, and Day of Week from a date column."""
    if date_col in df.columns and pd.api.types.is_datetime64_any_dtype(df[date_col]):
        df["Month"] = df[date_col].dt.strftime("%B")
        df["Quarter"] = "Q" + df[date_col].dt.quarter.astype(str)
        df["Day Of Week"] = df[date_col].dt.day_name()
        logging.info(f"Added derived date columns based on '{date_col}'")
    return df

def save_cleaned_csv(df: pd.DataFrame, output_path: str) -> None:
    """Save the DataFrame to a CSV file."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    df.to_csv(output_path, index=False)
    logging.info(f"Saved cleaned CSV to {output_path}")

# ──────────────────────────────────────────────────────────────
#  EXCEL REPORT GENERATION
# ──────────────────────────────────────────────────────────────

def _header_style(bg="1F3864", fg="FFFFFF", bold=True, size=11):
    fill = PatternFill("solid", fgColor=bg)
    font = Font(color=fg, bold=bold, size=size)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return fill, font, align

def _thin_border():
    thin = Side(style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def write_summary_sheet(ws, problems_before: Dict, df_before: pd.DataFrame, df_after: pd.DataFrame):
    ws.title = "Summary Dashboard"
    # ... (Keeping the styling identical to preserve the original logic, just condensed for space)
    
    def set_cell(r, c, value, bold=False, size=11, fg="595959", bg=None, center=False, merge_to=None, height=None):
        cell = ws.cell(row=r, column=c, value=value)
        cell.font = Font(bold=bold, size=size, color=fg)
        if bg: cell.fill = PatternFill("solid", fgColor=bg)
        if center: cell.alignment = Alignment(horizontal="center", vertical="center")
        if merge_to: ws.merge_cells(f"{get_column_letter(c)}{r}:{merge_to}{r}")
        if height: ws.row_dimensions[r].height = height
        return cell

    set_cell(1, 1, "DATA CLEANING & QUALITY REPORT", bold=True, size=18, fg="FFFFFF", bg="1F3864", center=True, merge_to="G", height=38)
    set_cell(2, 1, f"Generated on {datetime.now().strftime('%d %B %Y %H:%M')}", size=10, fg="A6A6A6", center=True, merge_to="G", height=18)
    
    # KPIs
    set_cell(4, 1, "DATASET OVERVIEW", bold=True, size=12, fg="FFFFFF", bg="2E74B5", center=True, merge_to="G", height=24)
    kpis = [
        ("Rows Before", len(df_before)), ("Rows After", len(df_after)),
        ("Columns", len(df_after.columns)), ("Duplicates Removed", problems_before.get("Duplicate rows", 0)),
        ("Blanks Fixed", sum(v for k, v in problems_before.items() if "Missing" in k and isinstance(v, int)))
    ]
    for i, (label, val) in enumerate(kpis, start=1):
        c = ws.cell(row=5, column=i, value=label)
        c.font, c.fill, c.alignment = Font(bold=True, size=9, color="FFFFFF"), PatternFill("solid", fgColor="2E74B5"), Alignment(horizontal="center", vertical="center", wrap_text=True)
        v = ws.cell(row=6, column=i, value=val)
        v.font, v.fill, v.alignment = Font(bold=True, size=20, color="1F3864"), PatternFill("solid", fgColor="D6E4F0"), Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[5].height, ws.row_dimensions[6].height = 30, 36

    # Issues
    set_cell(8, 1, "ISSUES FOUND & ACTIONS TAKEN", bold=True, size=12, fg="FFFFFF", bg="2E74B5", center=True, merge_to="G", height=24)
    for i, h in enumerate(["Issue", "Records Affected", "Action Taken"], start=1):
        c = ws.cell(row=9, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="595959"), Alignment(horizontal="center", vertical="center"), _thin_border()
    ws.merge_cells("C9:G9")
    
    row = 10
    for k, v in problems_before.items():
        bg = "D6E4F0" if row % 2 == 0 else "FFFFFF"
        action = "Removed — kept first occurrence" if "Duplicate" in k else ("Filled missing values" if "Missing" in k else "Standardised formatting")
        for col_i, val in enumerate([k, v, action], start=1):
            c = ws.cell(row=row, column=col_i, value=val)
            c.fill, c.border = PatternFill("solid", fgColor=bg), _thin_border()
            c.alignment = Alignment(vertical="center", horizontal="center" if col_i == 2 else "left", indent=0 if col_i == 2 else 1)
        ws.merge_cells(f"C{row}:G{row}")
        row += 1

    widths = {"A": 28, "B": 18, "C": 42, "D": 14, "E": 14, "F": 14, "G": 14}
    for col_letter, width in widths.items(): ws.column_dimensions[col_letter].width = width

def write_data_sheet(ws, df: pd.DataFrame):
    ws.title = "Cleaned Data"
    fill, font, align = _header_style()
    border = _thin_border()
    
    for col_idx, name in enumerate(df.columns, start=1):
        c = ws.cell(row=1, column=col_idx, value=name)
        c.fill, c.font, c.alignment, c.border = fill, font, align, border
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    
    for row_idx, row_data in enumerate(df.itertuples(index=False), start=2):
        bg = "F2F2F2" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, value in enumerate(row_data, start=1):
            if hasattr(value, "strftime"): value = value.strftime("%Y-%m-%d")
            elif value != value: value = ""
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            c.fill, c.border, c.alignment = PatternFill("solid", fgColor=bg), border, Alignment(vertical="center")

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 3, 28)
    ws.auto_filter.ref = ws.dimensions

def write_chart_sheet(ws, df: pd.DataFrame, category_col: str, value_col: str):
    ws.title = "Charts"
    if category_col not in df.columns or value_col not in df.columns:
        ws["A1"] = "Chart data not available (check column configurations)"
        return

    summary = df.groupby(category_col)[value_col].agg(["sum", "mean", "count"]).round(2).reset_index().sort_values("sum", ascending=False)
    summary.columns = [category_col, "Total", "Average", "Count"]

    ws["A1"] = "Summary by Category"
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    
    for i, h in enumerate(summary.columns, start=1):
        c = ws.cell(row=2, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864"), Alignment(horizontal="center", vertical="center"), _thin_border()
        
    for idx, row in enumerate(summary.itertuples(index=False), start=3):
        bg = "D6E4F0" if idx % 2 == 0 else "FFFFFF"
        for col_i, val in enumerate(row, start=1):
            c = ws.cell(row=idx, column=col_i, value=val)
            c.fill, c.border, c.alignment = PatternFill("solid", fgColor=bg), _thin_border(), Alignment(horizontal="left" if col_i == 1 else "center")
            
    chart = BarChart()
    chart.type, chart.style, chart.title, chart.grouping = "col", 10, f"Total {value_col} by {category_col}", "clustered"
    chart.width, chart.height = 18, 12
    data_ref = Reference(ws, min_col=2, min_row=2, max_row=2 + len(summary))
    cats_ref = Reference(ws, min_col=1, min_row=3, max_row=2 + len(summary))
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws.add_chart(chart, "F2")

def generate_excel_report(df_before: pd.DataFrame, df_after: pd.DataFrame, problems_before: Dict, output_path: str) -> None:
    logging.info("Generating Excel report...")
    if os.path.dirname(output_path):
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

    wb = Workbook()
    write_summary_sheet(wb.active, problems_before, df_before, df_after)
    write_data_sheet(wb.create_sheet(), df_after)
    write_chart_sheet(wb.create_sheet(), df_after, CATEGORY_COLUMN, VALUE_COLUMN)
    
    wb.save(output_path)
    logging.info(f"Successfully saved report to {output_path}")

# ──────────────────────────────────────────────────────────────
#  MAIN PIPELINE
# ──────────────────────────────────────────────────────────────

def run_pipeline() -> None:
    logging.info("Starting Data Cleaning Pipeline...")
    try:
        df = load_data(INPUT_FILE)
        df_original = df.copy()
        
        problems_before = inspect_data(df)
        
        df = fix_column_names(df)
        df = convert_numeric_columns(df, NUMERIC_COLUMNS)
        df = remove_duplicates(df)
        df = clean_text_columns(df, TEXT_COLUMNS)
        df = standardise_categories(df, CATEGORY_COLUMN)
        df = fill_missing_values(df, NUMERIC_COLUMNS, TEXT_COLUMNS)
        df = parse_dates(df, DATE_COLUMNS)
        df = add_derived_columns(df)
        df = flag_outliers(df, VALUE_COLUMN)
        
        save_cleaned_csv(df, OUTPUT_CSV)
        generate_excel_report(df_original, df, problems_before, OUTPUT_REPORT)
        
        logging.info("Pipeline completed successfully! ✅")
    except Exception as e:
        logging.error(f"Pipeline failed: {e}")

if __name__ == "__main__":
    run_pipeline()
